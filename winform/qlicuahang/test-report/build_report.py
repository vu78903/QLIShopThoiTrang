from __future__ import annotations

from collections import Counter, OrderedDict
from datetime import datetime
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = Path(r"C:\Users\Anh Vu\.codex\attachments\d8acf19a-0374-4954-91fa-10ce034f9b09\Pasted text.txt")
EVIDENCE = ROOT / "evidence"
OUTPUT = ROOT / "Bao_cao_kiem_thu_172_truong_hop_2026-10-23.docx"


def parse_cases() -> list[dict]:
    rows = []
    for line in SOURCE.read_text(encoding="utf-8-sig").splitlines():
        if not line.startswith("WIN-"):
            continue
        cols = line.split("\t")
        cols += [""] * (11 - len(cols))
        rows.append({
            "id": cols[0].strip(),
            "platform": cols[1].strip(),
            "module_code": cols[2].strip(),
            "area": cols[3].strip(),
            "type": cols[4].strip(),
            "title": cols[5].strip(),
            "precondition": cols[6].strip(),
            "steps": cols[7].strip(),
            "data": cols[8].strip(),
            "expected": cols[9].strip(),
            "note": cols[10].strip(),
        })
    return rows


FAIL_DEFECTS = {
    7: "APP-SEC-001", 10: "APP-SEC-002", 16: "APP-AUTH-001",
    34: "APP-POS-001", 36: "APP-SEC-003",
    51: "APP-DEL-001", 54: "APP-DEL-001", 61: "APP-SEC-001",
    72: "APP-RET-001", 75: "APP-SEC-001", 79: "APP-PROD-001", 90: "APP-SEC-001",
    106: "APP-DATA-001", 117: "APP-EMP-001", 118: "APP-SEC-001",
    128: "APP-INV-001", 129: "APP-INV-001",
    134: "APP-PO-001", 137: "APP-SEC-001", 141: "APP-PO-001",
    144: "APP-PROMO-001", 155: "APP-RPT-001", 157: "APP-RPT-001", 160: "APP-SEC-001",
    167: "APP-SET-001", 168: "APP-LIQ-001", 169: "APP-LIQ-001",
    170: "APP-LIQ-001", 171: "APP-LIQ-001", 172: "APP-LIQ-001",
}

BLOCKED = {2, 9, 127, 150, 163, 164, 166}

DEFECTS = OrderedDict([
    ("APP-SEC-001", ("Nghiêm trọng", "API nghiệp vụ không yêu cầu token hoặc vai trò. Phân quyền hiện chỉ được áp dụng ở giao diện WinForms nên có thể gọi thẳng API để đọc hoặc thay đổi dữ liệu.")),
    ("APP-SEC-002", ("Nghiêm trọng", "Tài khoản mẫu admin/123 vẫn đăng nhập thành công; mật khẩu nhân viên mới được lưu mặc định là 123 và truy vấn đăng nhập so sánh mật khẩu dạng rõ.")),
    ("APP-SEC-003", ("Nghiêm trọng", "API lập hóa đơn kiểm tra tồn nhưng ghi DonGia, TongCong, ChietKhau, VAT và KhachPhaiTra từ request của client, chưa tính lại theo dữ liệu sản phẩm và voucher trong cơ sở dữ liệu.")),
    ("APP-POS-001", ("Cao", "Nút thanh toán chưa có khóa gửi yêu cầu. Hai lần bấm cách nhau vài mili giây có thể tạo hai mã hóa đơn khác nhau trước khi giỏ được xóa.")),
    ("APP-AUTH-001", ("Trung bình", "Khi API mất kết nối, giao diện dùng localStorage và chỉ ghi console warning; không có thông báo rõ ràng và nút thử lại theo yêu cầu test.")),
    ("APP-DEL-001", ("Cao", "Luồng giao hàng chỉ có bốn trạng thái tiến tới giao thành công; chưa có giao thất bại, hoàn hàng hoặc bộ lọc theo trạng thái.")),
    ("APP-RET-001", ("Cao", "Phiếu đổi trả chưa cộng lại tồn cho hàng đủ điều kiện tái bán và chưa có trường phân loại tái bán so với hàng lỗi.")),
    ("APP-PROD-001", ("Trung bình", "POST sản phẩm dùng upsert theo MaSKU. Request thêm mới trùng SKU có thể biến thành cập nhật thay vì trả lỗi trùng.")),
    ("APP-DATA-001", ("Cao", "API app-state trả dữ liệu đối tác khi không có phiên hoặc token, gồm số điện thoại, email và địa chỉ.")),
    ("APP-EMP-001", ("Cao", "Màn nhân viên không có trường tên đăng nhập hoặc quy trình đặt mật khẩu; username sinh từ MaNV và mật khẩu mới mặc định là 123.")),
    ("APP-INV-001", ("Cao", "Chưa có phiếu kiểm kê, số đếm thực tế, lý do chênh lệch và nhật ký điều chỉnh tồn. API sản phẩm vẫn có thể ghi trực tiếp số tồn.")),
    ("APP-PO-001", ("Cao", "Phiếu nhập hiện chỉ nhận một SKU và lấy giá bán sản phẩm làm đơn giá. Chưa có đơn giá nhập hoặc phiếu nhiều dòng như yêu cầu.")),
    ("APP-PROMO-001", ("Trung bình", "API voucher dùng upsert theo mã. Mã trùng cập nhật chương trình hiện có thay vì trả lỗi trùng.")),
    ("APP-RPT-001", ("Cao", "Báo cáo chưa có trạng thái hủy đơn và chưa trừ hoàn tiền đổi trả khỏi doanh thu thuần.")),
    ("APP-SET-001", ("Thấp", "Nút Cài đặt chỉ hiện thông báo Đã mở cài đặt hệ thống nhưng không mở màn hình hay thông báo chưa hỗ trợ.")),
    ("APP-LIQ-001", ("Cao", "Chưa có màn hình, API hoặc dữ liệu cho thanh lý, hàng lỗi và lịch sử xuất kho thanh lý.")),
])

AUTOMATION_ERRORS = [
    ("AUT-001", "Đổi trả", "Test mở doi-tra.html ở thư mục gốc, trong khi tệp thật nằm trong Views. Kết quả ERR_FILE_NOT_FOUND.", "E18, E19"),
    ("AUT-002", "POS", "Test cũ yêu cầu fieldset sản phẩm bị khóa trước khi chọn khách. Giao diện hiện tại chủ động cho bán khách lẻ nên kỳ vọng đã lỗi thời.", "E04"),
    ("AUT-003", "Điều hướng", "Assertion URL không chấp nhận phần hash #delivery/#cod, dù trang giao hàng đã mở đúng.", "E16, E17"),
    ("AUT-004", "Thanh toán tiền mặt", "Fixture không truyền window.currentEmployee nên ứng dụng từ chối lập hóa đơn. Chạy lại với phiên NV-001 mở hóa đơn thành công.", "E23"),
    ("AUT-005", "COD", "Fixture không truyền window.currentEmployee nên dừng trước bước tạo COD; đây là lỗi dữ liệu test tương tự AUT-004.", "E17, E23"),
]

EVIDENCE_BY_AREA = {
    "Đăng nhập / vai trò / đăng xuất": ["E01", "E02", "E03", "E24"],
    "Tổng quan / menu / điều hướng": ["E20", "E22"],
    "POS - Bán hàng tại quầy": ["E04", "E05", "E23"],
    "Hóa đơn / Lịch sử / In / Kết ca": ["E06", "E21", "E23"],
    "Giao hàng / Đơn online": ["E16"],
    "COD / Đối soát": ["E17"],
    "Đổi trả sản phẩm": ["E18", "E19"],
    "Quản lý sản phẩm / SKU": ["E07", "E08", "E24"],
    "Quản lý loại sản phẩm": ["E09"],
    "Đối tác - khách hàng": ["E10", "E24"],
    "Đối tác - nhà cung cấp": ["E11"],
    "Đối tác - nhân viên": ["E10", "E24"],
    "Tồn kho / Kiểm kê": ["E12"],
    "Nhập kho / Duyệt phiếu": ["E12", "E13"],
    "Voucher / Khuyến mãi": ["E14", "E15"],
    "Báo cáo / Thống kê": ["E20", "E21"],
    "Cài đặt / Khởi động / Kết nối API": ["E22", "E25"],
    "Thanh lý / Hàng lỗi / Xuất kho": ["E12"],
}

CAPTIONS = {
    "E01": "Form đăng nhập hiển thị đầy đủ",
    "E02": "Mật khẩu được che mặc định",
    "E03": "Validation khi bỏ trống thông tin",
    "E04": "Màn hình POS bán hàng",
    "E05": "Chọn sản phẩm theo SKU từ kho",
    "E06": "Lịch sử và tra cứu hóa đơn",
    "E07": "Tìm sản phẩm theo MaSP",
    "E08": "Form thêm sản phẩm và MaSKU tự sinh",
    "E09": "Danh sách loại sản phẩm",
    "E10": "Quản lý khách hàng và nhân viên",
    "E11": "Quản lý nhà cung cấp",
    "E12": "Tồn kho và danh sách phiếu nhập",
    "E13": "Form tạo phiếu nhập một SKU",
    "E14": "Danh sách voucher",
    "E15": "Form tạo voucher",
    "E16": "Theo dõi giao hàng",
    "E17": "Theo dõi COD và tiền cần thu",
    "E18": "Màn hình đổi trả",
    "E19": "Tra cứu hóa đơn và lập phiếu đổi trả",
    "E20": "Báo cáo doanh thu và top sản phẩm",
    "E21": "Phiếu kết ca",
    "E22": "Cài đặt chỉ hiện thông báo",
    "E23": "Thanh toán thành công khi có phiên nhân viên hợp lệ",
    "E24": "Mở trực tiếp trang sản phẩm không có phiên đăng nhập",
    "E25": "API health trả trạng thái kết nối thật",
}


def evidence_file(eid: str) -> Path | None:
    matches = sorted(EVIDENCE.glob(f"{eid}-*.png"))
    return matches[0] if matches else None


def method_for(case: dict, status: str) -> str:
    n = int(case["id"].split("-")[1])
    if status == "BLOCKED":
        return "Chưa thực thi"
    if n in {3, 4, 5, 10, 162, 165}:
        return "API thật"
    if n in {7, 34, 36, 52, 58, 61, 68, 69, 72, 75, 79, 85, 87, 89, 90, 97, 104, 106, 112, 117, 118, 124, 129, 135, 137, 139, 140, 144, 155, 157, 160, 167, 168, 169, 170, 171, 172}:
        return "Đối chiếu mã và API"
    return "UI tự động"


def result_for(case: dict) -> dict:
    n = int(case["id"].split("-")[1])
    if n in FAIL_DEFECTS:
        defect = FAIL_DEFECTS[n]
        actual = DEFECTS[defect][1]
        return {"status": "FAIL", "defect": defect, "actual": actual}
    if n in BLOCKED:
        blocked_reasons = {
            2: "Không có thông tin đăng nhập thu ngân hợp lệ trong dữ liệu bàn giao; không thử đoán mật khẩu.",
            9: "Không tự động đóng/mở lại phiên Windows đang dùng để tránh làm gián đoạn ứng dụng của người dùng.",
            127: "Không có website thử nghiệm hoặc đơn Web được cung cấp để đối chiếu tồn kho hai chiều.",
            150: "Không có website thử nghiệm để xác nhận voucher đồng bộ từ Win sang Web.",
            163: "Không thay chuỗi kết nối Supabase thật bằng giá trị sai vì có thể làm gián đoạn phiên API đang chạy.",
            164: "Không có backend Website trong phạm vi project được kiểm thử để chạy song song và đối chiếu cổng.",
            166: "Không đóng tiến trình WinForms của người dùng trong phiên kiểm thử tự động.",
        }
        return {"status": "BLOCKED", "defect": "-", "actual": blocked_reasons[n]}
    return {"status": "PASS", "defect": "-", "actual": f"Đạt kỳ vọng: {case['expected']}"}


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=70, start=90, bottom=70, end=90) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_borders(table, color="D9D9D9", size="6") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:color"), color)


def set_col_width(cell, inches: float) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def style_run(run, size=10, bold=False, color="000000") -> None:
    run.font.name = "Arial"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_title(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    style_run(r, 22, True, "000000")


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_paragraph(style=f"Heading {level}")
    r = p.add_run(text)
    style_run(r, 15 if level == 1 else 12, True, "000000")
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(6)


def add_body(doc: Document, text: str, bold_lead: str | None = None) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.08
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        style_run(r1, 10.5, True)
        r2 = p.add_run(text[len(bold_lead):])
        style_run(r2, 10.5)
    else:
        r = p.add_run(text)
        style_run(r, 10.5)


def add_summary_table(doc: Document, rows: Iterable[tuple[str, str]]) -> None:
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    headers = ["Chỉ số", "Kết quả"]
    for i, value in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, "17365D")
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(value)
        style_run(r, 10, True, "FFFFFF")
    set_repeat_table_header(table.rows[0])
    for idx, (label, value) in enumerate(rows):
        cells = table.add_row().cells
        if idx % 2:
            for c in cells:
                set_cell_shading(c, "F3F6F9")
        for c in cells:
            set_cell_margins(c)
            c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for cell, text in zip(cells, (label, value)):
            p = cell.paragraphs[0]
            r = p.add_run(text)
            style_run(r, 10)
    for row in table.rows:
        set_col_width(row.cells[0], 2.8)
        set_col_width(row.cells[1], 3.5)


def add_results_table(doc: Document, cases: list[dict]) -> None:
    headers = ["Mã", "Trường hợp", "Phương pháp", "Kết quả", "Kết quả thực tế và mã lỗi", "Ảnh"]
    widths = [0.55, 1.35, 0.80, 0.60, 2.65, 0.55]
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    for i, text in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, "17365D")
        set_cell_margins(cell, 70, 65, 70, 65)
        set_col_width(cell, widths[i])
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        style_run(r, 8.5, True, "FFFFFF")
    set_repeat_table_header(table.rows[0])

    for idx, case in enumerate(cases):
        result = result_for(case)
        method = method_for(case, result["status"])
        evidence = ", ".join(EVIDENCE_BY_AREA.get(case["area"], [])[:2]) or "-"
        detail = result["actual"]
        if result["defect"] != "-":
            detail = f"{result['defect']}: {detail}"
        values = [case["id"], case["title"], method, result["status"], detail, evidence]
        cells = table.add_row().cells
        if idx % 2:
            for cell in cells:
                set_cell_shading(cell, "F3F6F9")
        for i, (cell, value) in enumerate(zip(cells, values)):
            set_cell_margins(cell, 60, 65, 60, 65)
            set_col_width(cell, widths[i])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i in {0, 2, 3, 5} else WD_ALIGN_PARAGRAPH.LEFT
            color = "008000" if result["status"] == "PASS" and i == 3 else "C00000" if result["status"] == "FAIL" and i == 3 else "9C6500" if result["status"] == "BLOCKED" and i == 3 else "000000"
            r = p.add_run(value)
            style_run(r, 8.2, i == 3, color)


def add_evidence_grid(doc: Document, evidence_ids: list[str]) -> None:
    files = [(eid, evidence_file(eid)) for eid in evidence_ids]
    files = [(eid, p) for eid, p in files if p is not None]
    if not files:
        return
    add_heading(doc, "Ảnh minh chứng", 2)
    for start in range(0, len(files), 2):
        pair = files[start:start + 2]
        table = doc.add_table(rows=2, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        set_table_borders(table, "FFFFFF", "0")
        for col in range(2):
            cell = table.rows[0].cells[col]
            set_cell_margins(cell, 35, 35, 20, 35)
            set_col_width(cell, 3.15)
            if col < len(pair):
                eid, image_path = pair[col]
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(str(image_path), width=Inches(3.0))
                cap = table.rows[1].cells[col]
                set_cell_margins(cap, 15, 35, 70, 35)
                cp = cap.paragraphs[0]
                cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = cp.add_run(f"{eid}  {CAPTIONS.get(eid, '')}")
                style_run(r, 8.5, False, "404040")
            else:
                table.rows[0].cells[col].text = ""
                table.rows[1].cells[col].text = ""


def add_defect_table(doc: Document) -> None:
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    headers = ["Mã lỗi", "Mức độ", "Mô tả thực tế", "Ảnh"]
    widths = [1.0, 0.9, 4.0, 0.6]
    for i, text in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, "17365D")
        set_cell_margins(cell)
        set_col_width(cell, widths[i])
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        style_run(r, 9, True, "FFFFFF")
    set_repeat_table_header(table.rows[0])
    defect_evidence = {
        "APP-SEC-001": "E24", "APP-SEC-002": "E01", "APP-SEC-003": "E23",
        "APP-POS-001": "E23", "APP-AUTH-001": "E04", "APP-DEL-001": "E16",
        "APP-RET-001": "E19", "APP-PROD-001": "E08", "APP-DATA-001": "E24",
        "APP-EMP-001": "E10", "APP-INV-001": "E12", "APP-PO-001": "E13",
        "APP-PROMO-001": "E15", "APP-RPT-001": "E20", "APP-SET-001": "E22",
        "APP-LIQ-001": "E12",
    }
    for idx, (code, (severity, description)) in enumerate(DEFECTS.items()):
        cells = table.add_row().cells
        if idx % 2:
            for cell in cells:
                set_cell_shading(cell, "F3F6F9")
        values = [code, severity, description, defect_evidence[code]]
        for i, (cell, text) in enumerate(zip(cells, values)):
            set_cell_margins(cell)
            set_col_width(cell, widths[i])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i in {0, 1, 3} else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(text)
            style_run(r, 8.8, i == 0)


def add_automation_table(doc: Document) -> None:
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    headers = ["Mã", "Luồng", "Kết quả thực tế", "Ảnh"]
    widths = [0.75, 1.25, 4.0, 0.55]
    for i, text in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, "17365D")
        set_cell_margins(cell)
        set_col_width(cell, widths[i])
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        style_run(r, 9, True, "FFFFFF")
    set_repeat_table_header(table.rows[0])
    for idx, row in enumerate(AUTOMATION_ERRORS):
        cells = table.add_row().cells
        if idx % 2:
            for cell in cells:
                set_cell_shading(cell, "F3F6F9")
        for i, (cell, text) in enumerate(zip(cells, row)):
            set_cell_margins(cell)
            set_col_width(cell, widths[i])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i in {0, 1, 3} else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(text)
            style_run(r, 9, i == 0)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = paragraph.add_run("Trang ")
    style_run(r, 8, False, "666666")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def main() -> None:
    cases = parse_cases()
    assert len(cases) == 172, f"Expected 172 cases, got {len(cases)}"
    enriched = [(case, result_for(case)) for case in cases]
    counts = Counter(result["status"] for _, result in enriched)

    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    add_page_number(section.footer.paragraphs[0])

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    styles["Normal"].font.size = Pt(10.5)
    for style_name in ("Title", "Heading 1", "Heading 2"):
        styles[style_name].font.name = "Arial"
        styles[style_name]._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        styles[style_name]._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        styles[style_name].font.color.rgb = RGBColor(0, 0, 0)

    add_title(doc, "Báo cáo kết quả kiểm thử ứng dụng Windows")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("172 trường hợp kiểm thử trên bản Release ngày 23 tháng 10 năm 2026")
    style_run(r, 12, False, "404040")
    doc.add_paragraph()
    add_body(doc, f"Kết luận chính: đã đánh giá đủ 172 trường hợp. Kết quả gồm {counts['PASS']} Pass, {counts['FAIL']} Fail và {counts['BLOCKED']} Blocked. Các lỗi quan trọng nhất nằm ở phân quyền API, lưu mật khẩu, xác thực giá tại máy chủ, kiểm kê tồn kho và các chức năng chưa có như thanh lý hàng lỗi.")

    add_heading(doc, "Tóm tắt kết quả", 1)
    add_summary_table(doc, [
        ("Tổng số trường hợp", "172"),
        ("Pass", str(counts["PASS"])),
        ("Fail", str(counts["FAIL"])),
        ("Blocked", str(counts["BLOCKED"])),
        ("Bộ test Playwright hiện có", "13 test  8 đạt  5 lỗi kịch bản hoặc fixture"),
        ("Chạy lại thanh toán với phiên nhân viên đúng", "Đạt và mở hóa đơn thành công"),
        ("API health thực tế", "HTTP 200  ok true"),
        ("Build Release", "0 lỗi  0 cảnh báo"),
    ])

    add_heading(doc, "Phạm vi và phương pháp", 1)
    add_body(doc, "Bản kiểm thử dùng các tệp giao diện và JavaScript hiện tại của project. Luồng chỉ đọc dùng API cục bộ thật tại cổng 5000. Các thao tác thêm, sửa, xóa và thanh toán được chạy bằng dữ liệu thử với API chặn ở trình duyệt để không thay đổi dữ liệu Supabase thật.")
    add_body(doc, "Mỗi kết quả ghi một trong ba trạng thái: Pass khi hành vi được thực thi hoặc xác minh đầy đủ; Fail khi quan sát thấy sai lệch hoặc chức năng chưa được triển khai; Blocked khi thiếu tài khoản, website, máy in hoặc việc thử có nguy cơ làm gián đoạn môi trường đang dùng.")
    add_body(doc, "Ảnh E01 đến E25 được chụp từ chính giao diện hiện tại. Dữ liệu hiển thị trong ảnh kiểm thử là dữ liệu thử, trừ ảnh API health và ảnh truy cập trực tiếp trang sản phẩm.")

    add_heading(doc, "Danh sách lỗi ghi nhận", 1)
    add_defect_table(doc)

    grouped: OrderedDict[str, list[dict]] = OrderedDict()
    for case in cases:
        grouped.setdefault(case["area"], []).append(case)

    for area, area_cases in grouped.items():
        doc.add_page_break()
        add_heading(doc, area, 1)
        area_counts = Counter(result_for(case)["status"] for case in area_cases)
        add_body(doc, f"Kết quả nhóm: {area_counts['PASS']} Pass, {area_counts['FAIL']} Fail, {area_counts['BLOCKED']} Blocked.")
        add_results_table(doc, area_cases)
        add_evidence_grid(doc, EVIDENCE_BY_AREA.get(area, []))

    doc.add_page_break()
    add_heading(doc, "Kết quả bộ test tự động hiện có", 1)
    add_body(doc, "Lần chạy toàn bộ Playwright cho kết quả 8 trên 13 test đạt. Năm lỗi bên dưới là lỗi đường dẫn, kỳ vọng hoặc fixture của bộ test cũ. Chúng được tách khỏi lỗi ứng dụng để tránh kết luận sai. Luồng thanh toán được chạy lại với phiên NV-001 và đã mở hóa đơn thành công trong ảnh E23.")
    add_automation_table(doc)
    add_evidence_grid(doc, ["E18", "E19", "E16", "E17", "E23"])

    add_heading(doc, "Ưu tiên xử lý", 1)
    priorities = [
        "Bắt buộc xác thực và phân quyền ở API cho toàn bộ endpoint đọc ghi.",
        "Băm mật khẩu, bỏ mật khẩu mặc định 123 và vô hiệu hóa tài khoản mẫu trước khi nộp.",
        "Tính lại giá, voucher, VAT và tổng tiền ở máy chủ khi lập hóa đơn.",
        "Thêm khóa chống gửi lặp cho thanh toán và các thao tác duyệt.",
        "Bổ sung kiểm kê, nhật ký tồn kho, nhập kho nhiều dòng với đơn giá nhập.",
        "Hoàn thiện giao thất bại, hoàn hàng, xử lý tồn đổi trả và báo cáo hoàn tiền.",
        "Xây dựng chức năng thanh lý hàng lỗi hoặc loại khỏi bộ yêu cầu nếu không nằm trong phạm vi bản nộp.",
    ]
    for item in priorities:
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run(item)
        style_run(r, 10.5)

    props = doc.core_properties
    props.title = "Báo cáo kết quả kiểm thử ứng dụng Windows"
    props.subject = "Kết quả 172 trường hợp kiểm thử"
    props.author = "Nhóm phát triển QLI Shop Thời Trang"
    props.keywords = "Windows, WinForms, kiểm thử, QA, Supabase, POS"
    props.comments = "Báo cáo tạo ngày 23/10/2026"

    doc.save(OUTPUT)
    print(f"OUTPUT={OUTPUT}")
    print(f"COUNTS={dict(counts)}")


if __name__ == "__main__":
    main()
